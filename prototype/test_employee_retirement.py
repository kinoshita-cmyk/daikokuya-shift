from __future__ import annotations

import ast
import copy
import tempfile
import unittest
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl

from prototype import employees
from prototype.employee_config import employee_from_dict, employee_to_dict
from prototype.excel_exporter import export_shift_to_excel
from prototype.export_config import export_column_order
from prototype.models import EmploymentStatus, MonthlyShift, ShiftAssignment, Store
from prototype.pdf_exporter import export_shift_to_pdf
from prototype.rule_book import build_month
from prototype.rule_config import CustomRule, RuleConfig
from prototype.shift_chat import PendingShiftChange, ShiftChatEngine
from prototype.submission_loader import SubmissionData
from prototype.validator import validate


class EmployeeRetirementTests(unittest.TestCase):
    def setUp(self):
        self.staff = copy.deepcopy(employees.ALL_EMPLOYEES)
        self.makino = next(e for e in self.staff if e.name == "牧野")
        employee_map = {e.name: e for e in self.staff}
        for target, value in (
            ("prototype.employees._dynamic_employee_map", employee_map),
            ("prototype.employee_config.get_all_employees_including_retired", self.staff),
        ):
            context = patch(target, return_value=value)
            context.start()
            self.addCleanup(context.stop)

    def test_default_and_restored_old_master_have_retirement_date(self):
        self.assertEqual(self.makino.retired_at, "2026-10-31")
        old = employee_to_dict(self.makino)
        old["retired_at"] = None
        restored = employee_from_dict(old)
        self.assertEqual(restored.retired_at, "2026-10-31")
        self.assertEqual(restored.affinities, self.makino.affinities)
        self.assertEqual(restored.full_name, self.makino.full_name)
        self.assertEqual(old["retired_at"], None)

    def test_october_roster_is_preserved_november_onwards_excludes(self):
        self.assertIn("牧野", {e.name for e in employees.shift_active_employees(2026, 10)})
        for year, month in ((2026, 11), (2026, 12), (2027, 1)):
            self.assertNotIn("牧野", {e.name for e in employees.shift_active_employees(year, month)})
        retired = copy.deepcopy(self.makino)
        retired.employment_status = EmploymentStatus.RETIRED
        self.assertTrue(employees.employee_is_shift_eligible(retired, 2026, 10))
        self.assertFalse(employees.employee_is_shift_eligible(retired, 2026, 11))

    def test_japan_date_boundary_and_status_label(self):
        with patch("prototype.submission_window.now_jst", return_value=datetime(2026, 10, 31, 23, 59)):
            self.assertTrue(employees.employee_is_in_service(self.makino))
            self.assertIn("退職予定", employees.employee_status_label(self.makino))
        with patch("prototype.submission_window.now_jst", return_value=datetime(2026, 11, 1)):
            self.assertFalse(employees.employee_is_in_service(self.makino))
            self.assertEqual(employees.employee_status_label(self.makino), "退職")

    def test_generator_uses_target_month_not_today(self):
        from prototype.generator import generate_shift
        with patch("prototype.generator.shift_active_employees", side_effect=RuntimeError("roster-check")) as roster:
            with self.assertRaisesRegex(RuntimeError, "roster-check"):
                generate_shift(2026, 11, {}, [], [], time_limit_seconds=1)
        roster.assert_called_once_with(2026, 11)

    def test_export_columns_preserve_recorded_history(self):
        self.assertIn("牧野", export_column_order(2026, 10))
        self.assertNotIn("牧野", export_column_order(2026, 11))
        assignment = ShiftAssignment("牧野", 1, Store.AKABANE)
        self.assertIn("牧野", export_column_order(2026, 11, [assignment]))

    def test_editor_uses_same_month_aware_roster(self):
        source = Path(__file__).resolve().parents[1] / "app" / "app.py"
        tree = ast.parse(source.read_text(encoding="utf-8"))
        node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "shift_to_editor_rows")
        namespace = {
            "MonthlyShift": MonthlyShift, "monthrange": monthrange, "date": date,
            "export_column_order": export_column_order,
            "detect_short_staff_by_store": lambda shift: {},
            "detect_key_warnings_by_store": lambda shift: {},
            "assignment_to_symbol": lambda assignment: "",
            "format_short_staff_summary_for_day": lambda stores: "",
            "format_key_warning_summary_for_day": lambda stores: "",
        }
        exec(compile(ast.Module(body=[node], type_ignores=[]), str(source), "exec"), namespace)
        october = namespace["shift_to_editor_rows"](MonthlyShift(2026, 10))
        november = namespace["shift_to_editor_rows"](MonthlyShift(2026, 11))
        self.assertIn("牧野", october[0])
        self.assertNotIn("牧野", november[0])
        self.assertEqual(len(november), 30)
        self.assertIn("人数少", november[0])
        self.assertIn("鍵", november[0])

    def test_excel_headers_and_right_column_widths_follow_retirement(self):
        with tempfile.TemporaryDirectory() as td:
            for month in (10, 11):
                path = Path(td) / f"{month}.xlsx"
                export_shift_to_excel(MonthlyShift(2026, month), path)
                workbook = openpyxl.load_workbook(path)
                sheet = workbook.active
                headers = [cell.value for cell in sheet[7]]
                self.assertEqual("牧野" in headers, month == 10)
                names = export_column_order(2026, month)
                short_col = 4 + len(names) + 2
                self.assertIn("人員", str(sheet.cell(7, short_col).value))
                self.assertEqual(sheet.column_dimensions[openpyxl.utils.get_column_letter(short_col)].width, 21)
                self.assertEqual(sheet.column_dimensions[openpyxl.utils.get_column_letter(short_col + 1)].width, 18)
                workbook.close()

    def test_pdf_receives_month_aware_columns(self):
        with tempfile.TemporaryDirectory() as td:
            for month in (10, 11):
                with patch("prototype.pdf_exporter.export_column_widths", wraps=__import__(
                    "prototype.excel_exporter", fromlist=["export_column_widths"]
                ).export_column_widths) as widths:
                    path = Path(td) / f"{month}.pdf"
                    export_shift_to_pdf(MonthlyShift(2026, month), path)
                self.assertEqual("牧野" in widths.call_args.args[0], month == 10)
                self.assertGreater(path.stat().st_size, 1000)

    def test_retirement_validation_does_not_hide_existing_work(self):
        assignment = ShiftAssignment("牧野", 1, Store.AKABANE)
        october = validate(MonthlyShift(2026, 10, assignments=[assignment]))
        november = validate(MonthlyShift(2026, 11, assignments=[assignment]))
        self.assertFalse(any(i.category == "退職後の勤務" for i in october.issues))
        errors = [i for i in november.issues if i.category == "退職後の勤務"]
        self.assertEqual([(i.employee, i.severity) for i in errors], [("牧野", "ERROR")])

    def test_ai_cannot_add_work_after_retirement_but_can_remove_it(self):
        engine = ShiftChatEngine(MonthlyShift(2026, 11), provider="local")
        self.assertTrue(engine._structural_change_violation_messages([
            PendingShiftChange(employee="牧野", day=1, store=Store.AKABANE),
        ]))
        for store in (Store.OFF, None):
            self.assertEqual(engine._structural_change_violation_messages([
                PendingShiftChange(employee="牧野", day=1, store=store),
            ]), [])
        engine.shift.month = 10
        self.assertEqual(engine._structural_change_violation_messages([
            PendingShiftChange(employee="牧野", day=1, store=Store.AKABANE),
        ]), [])

    def test_rule_book_excludes_future_staff_and_submissions_keeps_october(self):
        for month in (10, 11):
            sub = SubmissionData(2026, month, submitted_employees=["牧野"],
                                 off_requests={"牧野": [1]})
            book = build_month(2026, month, RuleConfig(), self.staff, {}, sub)
            self.assertEqual(any(e["name"] == "牧野" for e in book["staff"]), month == 10)
            self.assertEqual(any("牧野" in row["title"] for row in book["submissions"]), month == 10)

    def test_individual_rules_end_with_retirement_without_deleting_history(self):
        config = RuleConfig()
        config.custom_rules = [
            CustomRule("assign", "個人配置", "", employee="牧野",
                       rule_type="required_assignment", day=2, stores=["AKABANE"]),
            CustomRule("count", "巡回", "", employee="牧野",
                       rule_type="employee_store_count", count=1, stores=["AKABANE"]),
            CustomRule("note", "全体のメモ", "全体向け", rule_type="note"),
        ]
        source = Path(__file__).resolve().parents[1] / "app" / "app.py"
        tree = ast.parse(source.read_text(encoding="utf-8"))
        names = {
            "active_monthly_store_count_rules",
            "active_monthly_required_assignment_rules",
            "active_monthly_custom_rules",
        }
        nodes = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in names]
        namespace = {
            "RuleConfig": RuleConfig, "CustomRule": CustomRule,
            "system_monthly_store_count_rules": lambda y, m: [],
        }
        exec(compile(ast.Module(body=nodes, type_ignores=[]), str(source), "exec"), namespace)
        for name in names:
            october = namespace[name](config, 2026, 10)
            november = namespace[name](config, 2026, 11)
            self.assertGreater(len(october), len(november))
            if name != "active_monthly_custom_rules":
                self.assertEqual(november, [])
        self.assertEqual(len(config.custom_rules), 3)
        for month in (10, 11):
            book = build_month(2026, month, config, self.staff, {})
            ids = [c["id"] for c in book["individual"]]
            self.assertEqual("custom:assign" in ids, month == 10)
            self.assertEqual("custom:count" in ids, month == 10)
            self.assertIn("custom:note", ids)

    def test_existing_employee_session_expires_after_retirement(self):
        source = Path(__file__).resolve().parents[1] / "app" / "auth.py"
        tree = ast.parse(source.read_text(encoding="utf-8"))
        node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "require_auth")
        logout = Mock()
        ui = Mock(session_state={"authenticated": True})
        namespace = {
            "st": ui, "SESSION_AUTHENTICATED": "authenticated",
            "is_employee": lambda: True, "logout": logout,
            "get_logged_in_employee": lambda: "牧野",
            "get_employee": employees.get_employee,
            "employee_is_in_service": employees.employee_is_in_service,
        }
        exec(compile(ast.Module(body=[node], type_ignores=[]), str(source), "exec"), namespace)
        with patch("prototype.submission_window.now_jst", return_value=datetime(2026, 10, 31, 23, 59)):
            self.assertTrue(namespace["require_auth"]())
        logout.assert_not_called()
        with patch("prototype.submission_window.now_jst", return_value=datetime(2026, 11, 1)):
            self.assertFalse(namespace["require_auth"]())
        logout.assert_called_once()


if __name__ == "__main__":
    unittest.main()
